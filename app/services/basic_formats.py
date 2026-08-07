from __future__ import annotations

from pathlib import Path

from bs4 import BeautifulSoup
from charset_normalizer import from_path
from docx import Document
from ebooklib import ITEM_DOCUMENT, epub

from app.models import ExtractionResult


def process_text(path: Path, document_format: str) -> ExtractionResult:
    match = from_path(path).best()
    if match is None:
        raise ValueError("Unable to determine text encoding")
    return ExtractionResult(
        text=str(match),
        method=f"{document_format}_text",
        warnings=[f"Physical page count is unknown for {document_format.upper()}"],
    )


def process_docx(path: Path) -> ExtractionResult:
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text.strip() for cell in row.cells))
    properties = document.core_properties
    return ExtractionResult(
        text="\n\n".join(blocks),
        method="docx_text",
        title=properties.title or "",
        author=properties.author or "",
        warnings=["Physical page count is unknown for DOCX without rendering"],
    )


def process_epub(path: Path) -> ExtractionResult:
    book = epub.read_epub(str(path), options={"ignore_ncx": True})
    blocks: list[str] = []
    for item in book.get_items_of_type(ITEM_DOCUMENT):
        soup = BeautifulSoup(item.get_content(), "html.parser")
        text = soup.get_text("\n", strip=True)
        if text:
            blocks.append(text)
    title = next(iter(book.get_metadata("DC", "title")), ("", {}))[0]
    author = next(iter(book.get_metadata("DC", "creator")), ("", {}))[0]
    return ExtractionResult(
        text="\n\n".join(blocks),
        method="epub_text",
        title=title,
        author=author,
        warnings=["Physical page count is unknown for reflowable EPUB content"],
    )
